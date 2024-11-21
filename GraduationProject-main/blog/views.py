import io
import os
import csv
import re
import PyPDF2
from django.http import HttpResponse
from openpyxl import Workbook
from django.shortcuts import render
from docx import Document

# Function to extract name
def extract_name_from_text(text):
    """Extracts the name from the first line/paragraph of the text."""
    lines = text.split('\n')
    for line in lines:
        if line.strip():
            return line.strip()
    return None

# Function to extract emails
def extract_emails_from_text(text):
    """Extracts email addresses from the given text using a regex."""
    email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
    return re.findall(email_pattern, text)  # Returns a list of all emails found

# Function to extract age
def extract_age_from_text(text):
    """Extracts age if mentioned as a number followed by 'years' or 'age'."""
    age_pattern = r'\b(\d{1,2})\s*(?:years|year|yrs|age)\b'
    match = re.search(age_pattern, text, re.IGNORECASE)
    return match.group(1) if match else None

# Function to extract gender
def extract_gender_from_text(text):
    """Detects gender based on common keywords."""
    if re.search(r'\b(male|man|boy|he|him)\b', text, re.IGNORECASE):
        return "Male"
    elif re.search(r'\b(female|woman|girl|she|her)\b', text, re.IGNORECASE):
        return "Female"
    return "Not Specified"

# Function to extract education
def extract_education_from_text(text):
    """Extracts education based on keywords like degree names."""
    education_pattern = r'\b(Bachelor|Master|PhD|Diploma|Degree|High School|Certificate)\b'
    matches = re.findall(education_pattern, text, re.IGNORECASE)
    return ', '.join(set(matches)) if matches else "Not Specified"

# Main function to handle file uploads
def csv_to_excel(request):
    if request.method == 'POST' and 'csv_files' in request.FILES:
        csv_files = request.FILES.getlist('csv_files')  # Handle multiple files

        # Create a new Excel workbook
        workbook = Workbook()
        sheet = workbook.active

        # Add a header row
        sheet.append(["Name", "Email", "Age", "Gender", "Education"])

        for csv_file in csv_files:
            # Check file extension
            allowed_extensions = ['csv', 'tsv', 'xlsx', 'txt', 'pdf', 'docx']
            file_extension = csv_file.name.split('.')[-1].lower()
            if file_extension not in allowed_extensions:
                return HttpResponse(f"Unsupported file type: {csv_file.name}", status=400)

            try:
                # Extract text content from the file
                text = ""
                if file_extension in ['csv', 'tsv', 'txt']:
                    file_content = io.StringIO(csv_file.read().decode('utf-8'))
                    reader = csv.reader(file_content)
                    text = "\n".join([' '.join(row) for row in reader])
                elif file_extension == 'pdf':
                    pdf_content = io.BytesIO(csv_file.read())
                    reader = PyPDF2.PdfReader(pdf_content)
                    for page in reader.pages:
                        text += page.extract_text() + "\n"
                elif file_extension == 'docx':
                    doc_content = io.BytesIO(csv_file.read())
                    doc = Document(doc_content)
                    text = "\n".join([para.text for para in doc.paragraphs])

                # Extract data from text
                name = extract_name_from_text(text)
                emails = extract_emails_from_text(text)
                age = extract_age_from_text(text)
                gender = extract_gender_from_text(text)
                education = extract_education_from_text(text)

                # Write to Excel (add multiple rows if multiple emails exist)
                if emails:
                    for email in emails:
                        sheet.append([name, email, age, gender, education])
                else:
                    sheet.append([name, "No Email Found", age, gender, education])

            except Exception as e:
                return HttpResponse(f"Error reading file {csv_file.name}: {str(e)}", status=400)

        # Save the workbook to an external file
        output_file_path = os.path.join('output', 'extracted_data.xlsx')  # Save in an 'output' directory
        os.makedirs(os.path.dirname(output_file_path), exist_ok=True)
        workbook.save(output_file_path)

        return HttpResponse(f"Data has been extracted and saved to {output_file_path}. Please check the file.")

    return render(request, 'blog/upload_csv.html')
